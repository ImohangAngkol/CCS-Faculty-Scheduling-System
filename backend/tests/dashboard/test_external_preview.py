"""D3 real session authorization and memory-only parsing against synthetic metadata."""
import copy
import csv
from datetime import time
import hashlib
import io
import json
import random
import zipfile

import pytest
from fastapi import HTTPException
from sqlalchemy.orm import Session

from tests.dashboard.test_my_schedule import fixture, ORIGIN
from services import external_schedule_service as service, metadata_service
from routers import external_schedules
from schemas.external_schedule import ExternalRow

ROW = dict(supplier="Fictional College", subject_code="SYN-GEC101", title="Fictional Seminar",
           sections=["SYN-A"], academic_term="", days=["M"], start="09:00", end="10:00",
           room="SYN LEC", instructor="Synthetic Lecturer 0", source_reference="Fictional memo 1", notes="", status="tentative")
CTX = dict(sections=[dict(id="synthetic-section-a", label="SYN-A"), dict(id="synthetic-section-b", label="SYN-B")],
           rooms=[dict(id="SYN LEC", label="SYN LEC")],
           instructors=[dict(id="synthetic-instructor-0", label="Synthetic Lecturer 0", alias="Synthetic Faculty 0")],
           subjects=[dict(code="SYN101", title="Synthetic programming")], fixed=[])
REAL_CONTEXT = service.context


@pytest.fixture(autouse=True)
def synthetic_context(monkeypatch):
    monkeypatch.setattr(service, "context", lambda: copy.deepcopy(CTX))


def session(fixture, who="admin"):
    username = "admin@example.invalid" if who == "admin" else "faculty0@example.invalid"
    reply = fixture.client.post("/api/auth/login", headers={"Origin": ORIGIN}, json={"username": username, "password": fixture.password})
    assert reply.status_code == 200
    return {"Origin": ORIGIN, "X-CSRF-Token": reply.json()["csrf_token"]}


def csv_bytes(rows=(ROW,)):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream); writer.writerow(service.FIELDS)
    for row in rows:
        writer.writerow(["|".join(row.get(k, [])) if k in {"sections", "days"} else row.get(k, "") for k in service.FIELDS])
    return stream.getvalue().encode()


def xlsx_bytes(rows=(ROW,), mutate=None):
    from openpyxl import Workbook
    workbook = Workbook(); sheet = workbook.active
    sheet.append(list(service.FIELDS))
    for row in rows:
        sheet.append(["|".join(row.get(k, [])) if k in {"sections", "days"} else row.get(k, "") for k in service.FIELDS])
    if mutate:
        mutate(workbook, sheet)
    stream = io.BytesIO(); workbook.save(stream); workbook.close()
    return stream.getvalue()


def review(rows):
    return service.validate([ExternalRow(**r) for r in rows])


def codes(row):
    return {d.code for d in row.discrepancies}


@pytest.mark.parametrize("method,path", [("get","metadata"),("post","preview"),("post","preview-file?filename=x.csv")])
def test_anonymous_and_faculty_are_denied_before_body_or_metadata(fixture, monkeypatch, method, path):
    async def forbidden(_request):
        pytest.fail("Denied request must not read upload")
    monkeypatch.setattr(external_schedules, "bounded_body", forbidden)
    monkeypatch.setattr(service, "context", lambda: pytest.fail("Denied metadata access"))
    assert fixture.client.request(method, "/api/external-schedules/"+path).status_code == 401
    headers = session(fixture, "faculty")
    assert fixture.client.request(method, "/api/external-schedules/"+path, headers=headers).status_code == 403


def test_origin_and_csrf_are_required_even_for_non_persisted_post(fixture):
    headers = session(fixture)
    for bad in [{}, {"Origin":ORIGIN}, {**headers, "Origin":"https://evil.invalid"}]:
        response = fixture.client.post("/api/external-schedules/preview", headers=bad, json={"rows":[ROW]})
        assert response.status_code == 403
    assert fixture.client.post("/api/external-schedules/preview", headers=headers, json={"rows":[ROW]}).status_code == 200


def test_metadata_and_manual_contract_no_term_or_identity_fabrication(fixture):
    headers = session(fixture)
    data = fixture.client.get("/api/external-schedules/metadata").json()
    assert data["terms"] == [] and data["limits"]["max_rows"] == 200
    response = fixture.client.post("/api/external-schedules/preview", headers=headers, json={"rows":[ROW]})
    assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
    data = response.json()
    assert data["preview_only"] and not data["persisted"] and not data["activated"]
    assert data["source_sha256"] is None and data["source_name"] is None
    row = data["rows"][0]
    assert row["input"]["status"] == "tentative" and row["status"] == "needs_review"
    assert row["sections"][0]["id"] == "synthetic-section-a"
    assert "TERM_UNKNOWN" in {d["code"] for d in row["discrepancies"]}


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_valid_file_preview_all_rows_hashes_and_source_lines(fixture, kind):
    headers = session(fixture)
    raw = csv_bytes() if kind == "csv" else xlsx_bytes()
    response = fixture.client.post("/api/external-schedules/preview-file?filename=fictional."+kind, headers=headers, content=raw)
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["source_sha256"] == hashlib.sha256(raw).hexdigest() and data["source_kind"] == kind
    assert data["rows"][0]["source_row"] == 2 and data["rows"][0]["input"]["sections"] == ["SYN-A"]


@pytest.mark.parametrize("field,value,expected", [
    ("supplier","","REQUIRED_FIELD"),("sections",[],"REQUIRED_FIELD"),
    ("days",["THURSDAY","SUNDAY"],"INVALID_DAY"),("days",["M","M"],"INVALID_DAY"),
    ("start","9:00","INVALID_TIME"),("end","08:00","INVALID_TIME"),
    ("end","24:00","INVALID_TIME"),("start","09:10","UNSUPPORTED_TIME_GRID"),
    ("sections",["unknown"],"UNRESOLVED_SECTION"),("room","unknown","UNRESOLVED_ROOM"),
    ("instructor","","UNRESOLVED_INSTRUCTOR"),("academic_term","2026 imagined dates","TERM_UNVERIFIED"),
    ("sections",["SYN-A","syn-a"],"DUPLICATE_SECTION"),
])
def test_domain_errors_and_unresolved_references_remain_visible(field, value, expected):
    row = review([{**ROW, field:value}]).rows[0]
    assert expected in codes(row)
    if field in {"room","instructor"}:
        assert getattr(row,field).id is None


def test_duplicates_and_cross_resource_conflicts_are_symmetric_with_source_references():
    result = review([ROW, {**ROW,"sections":["syn-a"]}])
    for row in result.rows:
        assert {"DUPLICATE_ENTRY","SECTION_OVERLAP","ROOM_OVERLAP","INSTRUCTOR_OVERLAP"} <= codes(row)
        assert row.status == "invalid"
        assert all(d.related_rows for d in row.discrepancies if d.code.endswith("OVERLAP") or d.code=="DUPLICATE_ENTRY")


def test_adjacent_or_different_day_rows_do_not_overlap():
    result = review([ROW,{**ROW,"start":"10:00","end":"11:00"}, {**ROW,"days":["T"]}])
    assert all(not any("OVERLAP" in c for c in codes(row)) for row in result.rows)


def test_unknown_term_text_cannot_hide_a_possible_same_week_conflict():
    result = review([ROW,{**ROW,"academic_term":"Different unverified term"}])
    assert all("SECTION_OVERLAP" in codes(row) for row in result.rows)


def test_existing_fixed_comparison_is_read_only_and_labelled_provisional(monkeypatch):
    ctx = copy.deepcopy(CTX); ctx["fixed"]=[dict(section="SYN-A", day="M", start="08:30", end="09:30")]
    monkeypatch.setattr(service,"context",lambda:ctx)
    before = copy.deepcopy(ctx)
    assert "EXISTING_FIXED_SECTION_OVERLAP" in codes(review([ROW]).rows[0])
    assert ctx==before


def test_ambiguous_resource_name_never_picks_first_match(monkeypatch):
    ctx = copy.deepcopy(CTX); ctx["instructors"].append(dict(id="different",label=ROW["instructor"]))
    monkeypatch.setattr(service,"context",lambda:ctx)
    assert review([ROW]).rows[0].instructor.state=="unresolved"


@pytest.mark.parametrize("filename,raw,code", [
    ("a.xls", b"x", "UNSUPPORTED_FILE"),("../a.csv",b"x","INVALID_FILENAME"),
    ("a.csv",b"PK\x03\x04bad","FORMAT_MISMATCH"),("a.xlsx",b"csv text","FORMAT_MISMATCH"),
    ("a.csv",b"\xff","MALFORMED_CSV"),("a.csv",b"a,a\n1,2","INVALID_COLUMNS"),
    ("a.csv",b"","EMPTY_IMPORT"),("a.xlsx",b"PK\x03\x04broken","MALFORMED_XLSX"),
])
def test_file_rejections_are_atomic(filename, raw, code):
    with pytest.raises(HTTPException) as error:
        service.parse_file(raw, filename)
    assert error.value.detail["code"]==code and "rows" not in error.value.detail


@pytest.mark.parametrize("mutation,code", [
    (lambda w,s:s.cell(2,3,"=1+1"),"UNSAFE_WORKBOOK"),
    (lambda w,s:w.create_sheet("second"),"UNSAFE_WORKBOOK"),
    (lambda w,s:s.cell(2,3,123),"MALFORMED_ROW"),
    (lambda w,s:setattr(s.cell(2,3),"hyperlink","https://example.invalid"),"UNSAFE_WORKBOOK"),
    (lambda w,s:s.cell(500,1,"x"),"WORKSHEET_LIMIT"),
    (lambda w,s:s.cell(1,14,"x"),"WORKSHEET_LIMIT"),
])
def test_xlsx_formulas_links_multiple_sheets_types_and_bounds(mutation, code):
    with pytest.raises(HTTPException) as error:
        service.parse_file(xlsx_bytes(mutate=mutation),"fictional.xlsx")
    assert error.value.detail["code"]==code


def test_xlsx_time_only_values_are_supported():
    raw = xlsx_bytes(mutate=lambda w,s:(s.cell(2,7,time(9,0)),s.cell(2,8,time(10,0))))
    assert service.parse_file(raw,"fictional.xlsx").rows[0].input.start=="09:00"


@pytest.mark.parametrize("content", ["=IMPORTXML(\"x\")","+1+1","@SUM(A1)","-1+1","\x00bad"])
def test_csv_and_manual_formula_like_cells_reject_entire_batch(content):
    for action in [lambda:review([ROW,{**ROW,"notes":content}]),lambda:service.parse_file(csv_bytes([ROW,{**ROW,"notes":content}]),"x.csv")]:
        with pytest.raises(HTTPException):
            action()


@pytest.mark.parametrize("unsafe", ["../escape.xml","xl/vbaProject.bin","xl/externalLinks/externalLink1.xml","xl/embeddings/object.bin"])
def test_xlsx_disallowed_archive_members_are_rejected(unsafe):
    raw = xlsx_bytes()
    stream=io.BytesIO(raw)
    with zipfile.ZipFile(stream,"a") as archive:
        archive.writestr(unsafe,b"unsafe")
    with pytest.raises(HTTPException) as error:
        service.parse_file(stream.getvalue(),"x.xlsx")
    assert error.value.detail["code"]=="UNSAFE_WORKBOOK"


def test_archive_entity_declarations_rejected_before_xml_expansion():
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,"w") as archive:
        for name in ["[Content_Types].xml","_rels/.rels","xl/workbook.xml","xl/_rels/workbook.xml.rels","xl/worksheets/sheet1.xml"]:
            archive.writestr(name,b'<!DOCTYPE a [<!ENTITY x "bad">]><a>&x;</a>')
    with pytest.raises(HTTPException) as error:
        service.parse_file(stream.getvalue(),"x.xlsx")
    assert error.value.detail["code"]=="UNSAFE_WORKBOOK"


def test_zip_decompression_budget_rejects_a_small_compressed_bomb():
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,"w",compression=zipfile.ZIP_DEFLATED) as archive:
        for name in ["[Content_Types].xml","_rels/.rels","xl/workbook.xml","xl/_rels/workbook.xml.rels","xl/worksheets/sheet1.xml"]:
            archive.writestr(name,b"x"*(2*1024*1024))
    assert len(stream.getvalue()) < service.MAX_BYTES
    with pytest.raises(HTTPException) as error:
        service.parse_file(stream.getvalue(),"x.xlsx")
    assert error.value.detail["code"]=="UNSAFE_WORKBOOK"


def test_utf16_entity_payload_is_rejected_before_xml_decoding():
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,"w") as archive:
        for name in ["[Content_Types].xml","_rels/.rels","xl/workbook.xml","xl/_rels/workbook.xml.rels","xl/worksheets/sheet1.xml"]:
            archive.writestr(name,'<!DOCTYPE a [<!ENTITY x "bad">]><a>&x;</a>'.encode("utf-16"))
    with pytest.raises(HTTPException) as error:
        service.parse_file(stream.getvalue(),"x.xlsx")
    assert error.value.detail["code"]=="UNSAFE_WORKBOOK"


@pytest.mark.parametrize("xml", [b"<x>"*65+b"</x>"*65, b"<root>"+b"<x/>"*30001+b"</root>"], ids=["deep-xml", "many-nodes"])
def test_xml_depth_and_node_budgets_precede_workbook_loading(xml):
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,"w") as archive:
        for name in ["[Content_Types].xml","_rels/.rels","xl/workbook.xml","xl/_rels/workbook.xml.rels","xl/worksheets/sheet1.xml"]:
            archive.writestr(name,xml)
    with pytest.raises(HTTPException) as error:
        service.parse_file(stream.getvalue(),"x.xlsx")
    assert error.value.detail["code"]=="UNSAFE_WORKBOOK"


def test_streamed_body_without_content_length_is_still_bounded(fixture):
    headers=session(fixture)
    chunks=(b"x"*400000 for _ in range(3))
    response=fixture.client.post("/api/external-schedules/preview-file?filename=x.csv",headers=headers,content=chunks)
    assert response.status_code==413


def test_csv_multiline_notes_retain_physical_source_row_numbers():
    result=service.parse_file(csv_bytes([{**ROW,"notes":"Fictional line one\nFictional line two"},ROW]),"x.csv")
    assert [row.source_row for row in result.rows]==[2,4]


def test_upload_and_row_limits_and_no_partial_acceptance(fixture):
    headers=session(fixture)
    reply=fixture.client.post("/api/external-schedules/preview-file?filename=x.csv",headers=headers,content=b"x"*(service.MAX_BYTES+1))
    assert reply.status_code==413 and "rows" not in reply.json()["detail"]
    with pytest.raises(HTTPException) as error:
        service.parse_file(csv_bytes([ROW]*201),"x.csv")
    assert error.value.detail["code"]=="ROW_LIMIT"
    raw=csv_bytes()+b"too,few,columns\n"
    with pytest.raises(HTTPException) as error:
        service.parse_file(raw,"x.csv")
    assert error.value.detail["code"]=="MALFORMED_ROW"


def test_bad_manual_payload_duplicate_json_and_confirmed_status_rejected(fixture):
    headers={**session(fixture),"Content-Type":"application/json"}
    for raw in [b'{"rows":[]}',b'{"rows":[],"rows":[]}',json.dumps({"rows":[{**ROW,"status":"confirmed"}]}).encode(),json.dumps({"rows":[{**ROW,"faculty_id":"forged"}]}).encode()]:
        reply=fixture.client.post("/api/external-schedules/preview",headers=headers,content=raw)
        assert reply.status_code==422 and "rows" not in reply.json()["detail"]


def test_metadata_failure_is_retryable_and_sanitized(fixture,monkeypatch):
    session(fixture)
    def broken():
        raise ValueError("private path and private data")
    monkeypatch.setattr(service,"context",broken)
    reply=fixture.client.get("/api/external-schedules/metadata")
    assert reply.status_code==503 and "private" not in reply.text


def test_actual_metadata_adapter_uses_existing_stable_section_and_faculty_identities(fixture,monkeypatch):
    # Exercise the original adapter with synthetic metadata and keep isolation.
    from types import SimpleNamespace
    from services.identity_service import section_identity
    subject=SimpleNamespace(subject_code="SYN101",subject_title="Synthetic",offerings=[
        SimpleNamespace(section_id=section_identity("SYN-A"),section="SYN-A",preassigned_instructor=None)])
    monkeypatch.setattr(metadata_service,"list_subject_metadata",lambda:[subject])
    monkeypatch.setattr(metadata_service,"list_faculty_metadata",lambda:[])
    monkeypatch.setattr(metadata_service.scheduling,"list_subjects",fixture.courses)
    monkeypatch.setattr(metadata_service.scheduling,"lst_rooms",fixture.rooms)
    ctx=REAL_CONTEXT()
    assert ctx["sections"][0]["id"]==section_identity("SYN-A")


def test_preview_preserves_session_rows_storage_and_rng(fixture):
    headers=session(fixture)
    def db_digest():
        with fixture.engine.connect() as db:
            return repr([db.exec_driver_sql("SELECT * FROM "+table+" ORDER BY id").all() for table in ["users","login_sessions","faculty_identities"]])
    before=db_digest(); files={p:p.read_bytes() for p in fixture.path.parent.iterdir() if p.is_file()}; rng=random.getstate()
    first=fixture.client.post("/api/external-schedules/preview",headers=headers,json={"rows":[ROW]}).json()
    second=fixture.client.post("/api/external-schedules/preview",headers=headers,json={"rows":[ROW]}).json()
    assert first==second and first["preview_only"]
    assert db_digest()==before and random.getstate()==rng
    assert all(path.read_bytes()==value for path,value in files.items())
    assert set(fixture.path.parent.iterdir())==set(files)
