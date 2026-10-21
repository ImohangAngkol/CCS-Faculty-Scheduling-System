"""Bounded, memory-only external schedule parsing and review. No GA activation."""
import csv
import hashlib
import io
import re
import zipfile
from datetime import time
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET

from fastapi import HTTPException
from pydantic import ValidationError

from schemas.external_schedule import ExternalRow, PreviewResult, ReviewedRow, ReferenceMatch, Discrepancy

MAX_BYTES = 1024 * 1024
MAX_ROWS = 200
MAX_XML_BYTES = 8 * 1024 * 1024
MAX_ISSUES = 5000
FIELDS = tuple(ExternalRow.model_fields)
REQUIRED = {"supplier", "subject_code", "title", "sections", "days", "start", "end", "source_reference"}
DAYS = ("M", "T", "W", "TH", "F", "S")
DAY_ALIASES = dict(zip(("MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY"), DAYS))


def fail(code, message, status=422):
    raise HTTPException(status, {"code": code, "message": message, "preview_only": True})


def key(value):
    return str(value).strip().casefold()


def context():
    """Read existing metadata/resources; term registry is deliberately unconfigured."""
    from services import metadata_service as metadata
    from services.identity_service import section_identity
    scheduling = metadata.scheduling
    sections, instructors, subjects, fixed = {}, {}, [], []
    for subject in metadata.list_subject_metadata():
        subjects.append({"code": subject.subject_code, "title": subject.subject_title})
        for offering in subject.offerings:
            sections[offering.section_id] = {"id": offering.section_id, "label": offering.section}
            if offering.preassigned_instructor:
                item = offering.preassigned_instructor
                instructors[item.instructor_id] = {"id": item.instructor_id, "label": item.display_name}
    for member in metadata.list_faculty_metadata():
        instructors[member.faculty_id] = {"id": member.faculty_id, "label": member.name, "alias": member.display_code}
    seen = set()
    for subject in scheduling.list_subjects:
        identifier = section_identity(subject.section.code)
        if identifier not in seen:
            seen.add(identifier)
            fixed.extend(scheduling.get_fixed_section_schedule_entries(subject.section))
    rooms = [{"id": room.name, "label": room.name} for room in scheduling.lst_rooms]
    return {"sections": sorted(sections.values(), key=lambda r: r["label"]),
            "rooms": sorted(rooms, key=lambda r: r["label"]),
            "instructors": sorted(instructors.values(), key=lambda r: r["label"]),
            "subjects": subjects, "fixed": fixed}


def safe_context():
    try:
        return context()
    except Exception:
        fail("REFERENCE_METADATA_UNAVAILABLE", "Reference metadata is unavailable. Retry without changing scheduling inputs.", 503)


def metadata():
    ctx = safe_context()
    return {name: ctx[name] for name in ("sections", "rooms", "instructors", "subjects")} | {
        "terms": [], "term_verification": "not_configured",
        "fixed_section_entries": len(ctx["fixed"]),
        "limits": {"max_bytes": MAX_BYTES, "max_rows": MAX_ROWS, "max_columns": len(FIELDS),
                   "xlsx_sheets": 1, "max_decompressed_bytes": MAX_XML_BYTES},
    }


def _match(value, choices):
    if not value.strip():
        return ReferenceMatch(value=value, id=None, state="not_supplied")
    matches = {r["id"] for r in choices if key(value) in {key(r["id"]), key(r["label"]), key(r.get("alias", ""))}}
    return ReferenceMatch(value=value, id=next(iter(matches)) if len(matches) == 1 else None,
                          state="resolved" if len(matches) == 1 else "unresolved")


def _minutes(value):
    if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", value):
        return None
    h, m = map(int, value.split(":"))
    return h * 60 + m


def _unsafe_text(value):
    return any(ord(c) < 32 and c not in "\n\r\t" for c in value) or value.lstrip().startswith(("=", "+", "-", "@"))


def validate(rows, source_kind="manual", source_name=None, raw=None, source_rows=None, *, _context=None, _sign=True):
    ctx = safe_context() if _context is None else _context
    output, intervals, duplicates = [], [], {}
    issue_count = 0

    def issue(target, code, severity, field, message, related=()):
        nonlocal issue_count
        issue_count += 1
        if issue_count > MAX_ISSUES:
            fail("TOO_MANY_DISCREPANCIES", "This batch has too many conflicts to review safely. Split it into smaller previews.")
        target.discrepancies.append(Discrepancy(code=code, severity=severity, field=field, message=message, related_rows=list(related)))

    for index, original in enumerate(rows):
        values = original.model_dump()
        for name, value in values.items():
            texts = value if isinstance(value, list) else [value]
            if any(not isinstance(text, str) or len(text) > 1000 or _unsafe_text(text) for text in texts):
                fail("UNSAFE_CELL", "A field contains formula-like text, control characters or an oversized value. No rows were accepted.")
            values[name] = [text.strip() for text in value] if isinstance(value, list) else value.strip()
        values["days"] = [DAY_ALIASES.get(day.upper(), day.upper()) for day in values["days"]]
        row = ExternalRow(**values)
        line = source_rows[index] if source_rows else index + 1
        result = ReviewedRow(source_row=line, input=row, status=row.status,
                             sections=[_match(v, ctx["sections"]) for v in row.sections],
                             room=_match(row.room, ctx["rooms"]), instructor=_match(row.instructor, ctx["instructors"]),
                             discrepancies=[])
        output.append(result)
        for field in REQUIRED:
            if not getattr(row, field):
                issue(result, "REQUIRED_FIELD", "error", field, "This required field is missing.")
        if any(not v for v in row.sections):
            issue(result, "REQUIRED_FIELD", "error", "sections", "Section references must not be blank.")
        if len(set(map(key, row.sections))) != len(row.sections):
            issue(result, "DUPLICATE_SECTION", "error", "sections", "The same section is listed more than once.")
        if any(day not in DAYS for day in row.days) or len(set(row.days)) != len(row.days):
            issue(result, "INVALID_DAY", "error", "days", "Use distinct M, T, W, TH, F or S days. Sunday is unsupported.")
        start, end = _minutes(row.start), _minutes(row.end)
        valid_time = start is not None and end is not None and start < end
        if not valid_time:
            issue(result, "INVALID_TIME", "error", "start/end", "Use HH:MM in 24-hour time with start before end; overnight meetings are unsupported.")
        elif start % 30 or end % 30 or start < 450 or end > 1320:
            issue(result, "UNSUPPORTED_TIME_GRID", "review", "start/end", "This interval is outside the existing 07:30–22:00 half-hour grid. No times were rounded.")
        issue(result, "TERM_UNVERIFIED" if row.academic_term else "TERM_UNKNOWN", "review", "academic_term",
              "No verified academic term/date registry is configured. Dates and cross-term conflict coverage cannot be established.")
        if not any(key(row.subject_code) == key(s["code"]) for s in ctx["subjects"]):
            issue(result, "SUBJECT_NOT_IN_CATALOG", "review", "subject_code", "External subject is not resolved in the current prototype catalog.")
        for field, matches in (("sections", result.sections), ("room", [result.room]), ("instructor", [result.instructor])):
            if any(match.state != "resolved" for match in matches):
                issue(result, "UNRESOLVED_" + field.upper().rstrip("S"), "review", field,
                      "Reference is missing, unknown or ambiguous. No identity was created; resource conflict coverage is incomplete.")
        duplicate = (key(row.supplier), key(row.subject_code), tuple(sorted(map(key, row.sections))),
                     tuple(sorted(row.days)), row.start, row.end, key(row.room), key(row.instructor), key(row.academic_term))
        if duplicate in duplicates:
            prior = duplicates[duplicate]
            issue(result, "DUPLICATE_ENTRY", "error", "row", "Duplicate meeting within this preview.", [prior.source_row])
            issue(prior, "DUPLICATE_ENTRY", "error", "row", "Duplicate meeting within this preview.", [line])
        else:
            duplicates[duplicate] = result
        if valid_time and row.days and all(day in DAYS for day in row.days):
            intervals.append((result, start, end))
            for fixed in ctx["fixed"]:
                a, b = _minutes(fixed["start"]), _minutes(fixed["end"])
                if (key(fixed["section"]) in set(map(key, row.sections)) and fixed["day"] in row.days
                        and a is not None and b is not None and start < b and a < end):
                    issue(result, "EXISTING_FIXED_SECTION_OVERLAP", "review", "sections",
                          "Potential overlap with a loaded legacy fixed-section interval. Its term/dates are unverified; no existing block was changed.")
    for index, (row, start, end) in enumerate(intervals):
        for other, a, b in intervals[:index]:
            if not set(row.input.days) & set(other.input.days) or not (start < b and a < end):
                continue
            clashes = []
            if set(map(key, row.input.sections)) & set(map(key, other.input.sections)):
                clashes.append(("SECTION_OVERLAP", "sections"))
            for field in ("room", "instructor"):
                left, right = getattr(row, field), getattr(other, field)
                if left.id is not None and left.id == right.id:
                    clashes.append((field.upper() + "_OVERLAP", field))
            for code, field in clashes:
                issue(row, code, "error", field, "Potential same-week resource overlap within this batch; academic dates remain unverified.", [other.source_row])
                issue(other, code, "error", field, "Potential same-week resource overlap within this batch; academic dates remain unverified.", [row.source_row])
    for row in output:
        row.status = "invalid" if any(d.severity == "error" for d in row.discrepancies) else "needs_review" if row.discrepancies else row.input.status
    result = PreviewResult(source_kind=source_kind, source_name=source_name,
                           source_sha256=hashlib.sha256(raw).hexdigest() if raw is not None else None, rows=output)
    if _sign:
        from services.external_schedule_review_service import sign_preview
        result.review_token = sign_preview(result, ctx)
    return result


def filename_kind(filename):
    if (not filename or len(filename) > 120 or "/" in filename or "\\" in filename
            or any(ord(c) < 32 for c in filename)):
        fail("INVALID_FILENAME", "Provide a plain filename without a directory path.")
    extension = PurePosixPath(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        fail("UNSUPPORTED_FILE", "Only UTF-8 comma-separated CSV and bounded single-sheet XLSX are supported.")
    return extension[1:]


def _records(records):
    try:
        header_line, header = next(records)
    except StopIteration:
        fail("EMPTY_IMPORT", "The file has no header or schedule rows.")
    if (not header or len(header) > len(FIELDS) or any(not isinstance(c, str) for c in header)
            or len(set(header)) != len(header) or set(header) - set(FIELDS) or not REQUIRED <= set(header)):
        fail("INVALID_COLUMNS", "Use unique canonical template columns, including all required columns. Unknown columns are unsupported.")
    rows, lines = [], []
    try:
        for line, cells in records:
            if len(rows) >= MAX_ROWS:
                fail("ROW_LIMIT", "Preview is limited to 200 rows. No partial preview was accepted.")
            if len(cells) != len(header) or any(not isinstance(c, str) for c in cells):
                fail("MALFORMED_ROW", "Every source row must match the header and use text cells (or time-only XLSX start/end cells).")
            values = dict(zip(header, cells))
            for name in ("sections", "days"):
                values[name] = values.get(name, "").split("|") if values.get(name) else []
            rows.append(ExternalRow(**values)); lines.append(line)
    except ValidationError:
        fail("INVALID_ROW_STRUCTURE", "A source row exceeds field limits or contains unsupported status/types. No rows were accepted.")
    if not rows:
        fail("EMPTY_IMPORT", "The file contains no schedule rows.")
    return rows, lines


def _csv(raw):
    if raw.startswith((b"PK", b"\xd0\xcf\x11\xe0")):
        fail("FORMAT_MISMATCH", "File contents do not match UTF-8 CSV.")
    try:
        text = raw.decode("utf-8-sig")
        if "\x00" in text:
            raise ValueError()
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        def records():
            previous = 0
            for cells in reader:
                yield previous + 1, cells
                previous = reader.line_num
        return _records(iter(records()))
    except (UnicodeError, csv.Error, ValueError):
        fail("MALFORMED_CSV", "CSV must be valid UTF-8 with comma separators and correctly quoted fields.")


def _xlsx(raw):
    allowed = {"[Content_Types].xml", "_rels/.rels", "docProps/app.xml", "docProps/core.xml",
               "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml", "xl/sharedStrings.xml",
               "xl/theme/theme1.xml", "xl/worksheets/sheet1.xml"}
    required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/worksheets/sheet1.xml"}
    if not raw.startswith(b"PK\x03\x04"):
        fail("FORMAT_MISMATCH", "File contents do not match an XLSX workbook.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            xml_nodes = 0
            infos = archive.infolist(); names = [i.filename for i in infos]
            if (len(infos) > 32 or len(set(names)) != len(names) or set(names) - allowed or not required <= set(names)
                    or sum(i.file_size for i in infos) > MAX_XML_BYTES
                    or any(i.flag_bits & 1 or i.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
                           or i.file_size > 4*1024*1024 or i.file_size > max(i.compress_size, 1)*1000 for i in infos)):
                fail("UNSAFE_WORKBOOK", "Workbook exceeds archive limits or contains unsupported sheets, macros, links or embedded content.")
            for info in infos:
                data = archive.read(info)
                if b"\x00" in data or b"<!doctype" in data.lower() or b"<!entity" in data.lower():
                    fail("UNSAFE_WORKBOOK", "XML declarations/entities are not supported.")
                depth = 0
                for event, node in ET.iterparse(io.BytesIO(data), events=("start", "end")):
                    if event == "end":
                        depth -= 1
                        node.clear()
                        continue
                    depth += 1
                    xml_nodes += 1
                    if depth > 64 or xml_nodes > 30000:
                        fail("UNSAFE_WORKBOOK", "Workbook XML exceeds depth or node-count limits.")
                    tag = node.tag.split("}")[-1]
                    if tag in {"f", "definedName", "oleObject", "externalReference", "hyperlink"} or node.attrib.get("TargetMode") == "External":
                        fail("UNSAFE_WORKBOOK", "Formulas, defined names, external links and active workbook content are unsupported.")
                    if tag == "Relationship":
                        target = node.attrib.get("Target", "")
                        if not target or "\\" in target or ":" in target or ".." in PurePosixPath(target).parts:
                            fail("UNSAFE_WORKBOOK", "Unsafe workbook relationship target.")
                        relative = "" if target.startswith("/") or info.filename == "_rels/.rels" else "xl/"
                        if relative + target.lstrip("/") not in allowed:
                            fail("UNSAFE_WORKBOOK", "Unsupported workbook relationship target.")
                    if info.filename == "xl/worksheets/sheet1.xml" and tag == "c":
                        cell = node.attrib.get("r", "")
                        if not re.fullmatch(r"[A-M][1-9]\d{0,2}", cell) or int(re.search(r"\d+", cell).group()) > MAX_ROWS+1:
                            fail("WORKSHEET_LIMIT", "Worksheet must fit 13 columns and 200 data rows.")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                fail("UNSAFE_WORKBOOK", "Use one worksheet only.")
            sheet = workbook.worksheets[0]
            if sheet.max_row is None or sheet.max_column is None or sheet.max_row > MAX_ROWS+1 or sheet.max_column > len(FIELDS):
                fail("WORKSHEET_LIMIT", "Worksheet dimensions exceed preview limits.")
            def records():
                header = None
                for line, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = []
                    for index, value in enumerate(values):
                        name = header[index] if header and index < len(header) else None
                        if isinstance(value, time) and name in {"start", "end"} and not value.second and not value.microsecond:
                            value = value.strftime("%H:%M")
                        cells.append("" if value is None else value)
                    if line == 1:
                        header = cells
                    yield line, cells
            return _records(iter(records()))
        finally:
            workbook.close()
    except HTTPException:
        raise
    except Exception:
        fail("MALFORMED_XLSX", "The workbook is malformed or unsupported. No rows were accepted.")


def parse_file(raw, filename):
    kind = filename_kind(filename)
    if len(raw) > MAX_BYTES:
        fail("FILE_TOO_LARGE", "Preview uploads are limited to 1 MiB.", 413)
    rows, lines = _csv(raw) if kind == "csv" else _xlsx(raw)
    return validate(rows, kind, filename, raw, lines)
